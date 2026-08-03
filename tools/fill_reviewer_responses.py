from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "jimaging-4429747" / "MDPI_Article_Template"
OUT_DIR = TEMPLATE_DIR / "response_letters"

SOURCE = {
    1: Path("/Users/songyuexin/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/wxid_ywa1tvfadxtq22_628f/temp/drag/Response to reviewer 1 - MDPI(1).docx"),
    2: Path("/Users/songyuexin/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/wxid_ywa1tvfadxtq22_628f/temp/drag/Response to reviewer 2 - MDPI(1).docx"),
    3: Path("/Users/songyuexin/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/wxid_ywa1tvfadxtq22_628f/temp/drag/Response to reviewer 3 - MDPI.docx"),
    4: Path("/Users/songyuexin/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/wxid_ywa1tvfadxtq22_628f/temp/drag/Response to reviewer 4 - MDPI.docx"),
}


SUMMARY = (
    "We thank the reviewers for their careful and constructive assessment. We revised the manuscript in five coordinated ways: "
    "we narrowed the novelty claim to detector-internal reliability propagation rather than claiming that FFT, contrastive learning, "
    "or uncertainty regression is new in isolation; we added a compact ExDark component ablation and additional Mine-Objects controls; "
    "we clarified dataset partitions, baseline protocols, symbols, equations, and metric definitions; we synchronized the quantitative and "
    "qualitative figures with the final validation and test reporting; and we expanded the efficiency, statistical, and dataset limitations. "
    "We also provide the exact ScienceDB record and public source-code links. The revised manuscript avoids claims of zero-shot transfer, "
    "universal backbone independence, or conclusive uncertainty calibration."
)


GENERAL = {
    1: [
        ("The introduction provide sufficient background and include all relevant references?", "Can be improved", "We expanded the related-work discussion to include recent degraded-scene and frequency-aware DETR methods and clarified the distinction between established operators and our cross-stage reliability interface (Section 2)."),
        ("Is the research design appropriate?", "Yes", "The controlled RF-DETR protocol is retained, with a fixed split, common backbone, training budget, post-processing, and validation-based checkpoint selection for the module study (Sections 4.1 and 4.2)."),
        ("Are the methods adequately described?", "Yes", "We added definitions for all variables requested by the reviewer, explained FiLM and the real FFT, and stated the final module settings and training-only paths (Sections 3.1--3.4)."),
        ("Are the results clearly presented?", "Yes", "We added ExDark component ablation, coupling controls, band-count sensitivity, SCU coefficient sensitivity, and degradation-stratified validation analysis (Tables 4--6 and Figures 3--4)."),
        ("Are the conclusions supported by the results?", "Yes", "The conclusion now describes the method as a reliability framework and explicitly limits the interpretation to dataset-specific evaluation and supporting diagnostics."),
        ("Are all figures and tables clear and well-presented?", "Yes", "Figure 5 now explicitly states that each column compares the same input, RF-DETR baseline, and Spectral-DETR output; the quantitative figures were synchronized with the final tables."),
    ],
    2: [
        ("The introduction provide sufficient background and include all relevant references?", "Yes", "The introduction and related work were retained after a final clarity pass; the requested method definitions and recent references are now anchored where they are used."),
        ("Is the research design appropriate?", "Yes", "The evaluation distinguishes validation-based ablations from held-out test comparisons and does not present ScienceDB or ExDark as zero-shot transfer."),
        ("Are the methods adequately described?", "Yes", "The notation and implementation details were expanded throughout Section 3 and the evaluation protocol."),
        ("Are the results clearly presented?", "Yes", "The revised tables and figures identify datasets, partitions, metrics, and module configurations directly in captions or surrounding text."),
        ("Are the conclusions supported by the results?", "Yes", "Claims were narrowed where the evidence is diagnostic rather than causal, especially for gate statistics and uncertainty correlation."),
        ("Are all figures and tables clear and well-presented?", "Yes", "Figure 2 is accompanied by explicit FiLM and FFT definitions, and Figure 5 uses matched row-wise baseline comparisons."),
    ],
    3: [
        ("The introduction provide sufficient background and include all relevant references?", "Yes", "The related-work section now positions Spectral-DETR against recent low-light, degraded-scene, and frequency-aware detector variants."),
        ("Is the research design appropriate?", "Yes", "The revision keeps the fixed Mine-Objects split and adds focused controls rather than changing the central experimental protocol."),
        ("Are the methods adequately described?", "Yes", "The revised method section defines the band parameters, gate statistic, temperature, salience target, and training-only objectives."),
        ("Are the results clearly presented?", "Yes", "The revised presentation includes band sensitivity, SCU sensitivity, coupling controls, ExDark ablation, and the ScienceDB localization-gap analysis."),
        ("Are the conclusions supported by the results?", "Yes", "We explicitly describe the interaction as modest and the uncertainty correlation as supporting evidence, not proof of a calibrated physical severity estimator."),
        ("Are all figures and tables clear and well-presented?", "Yes", "Figures 1--5 and their captions were revised for module paths, definitions, matched qualitative comparisons, and synchronization with the final tables."),
    ],
    4: [
        ("The introduction provide sufficient background and include all relevant references?", "Yes", "The review is now written as connected paragraphs and includes recent detector and frequency-domain work."),
        ("Is the research design appropriate?", "Yes", "The fixed sequence-level Mine-Objects split and validation/test roles are stated explicitly."),
        ("Are the methods adequately described?", "Yes", "The revised text defines the requested symbols and gives the final settings needed to reproduce the reported protocol."),
        ("Are the results clearly presented?", "Yes", "The tables, ablations, diagnostics, and limitations were synchronized with the revised narrative."),
        ("Are the conclusions supported by the results?", "Yes", "The conclusion avoids universal generalization and presents the three stages as established operators connected by a detector-specific reliability pathway."),
        ("Are all figures and tables clear and well-presented?", "Yes", "The final compiled manuscript will be checked after the updated figure files are included; figure captions now identify panel roles and comparison arrangements explicitly."),
    ],
}


COMMENTS = {
1: [
 ("The three modules are combinations of existing techniques and lack theoretical breakthrough innovation.", "We agree that FFT filtering, contrastive learning, and uncertainty regression are established ingredients. We therefore revised the positioning throughout the Introduction and Related Work: Spectral-DETR no longer claims that any individual operator is theoretically novel. The contribution is framed more narrowly as a detector-internal reliability pathway in which learnable frequency gates condition query-level temperature and SCU-calibrated uncertainty weights localization supervision. This distinction, and the boundary of the novelty claim, are stated in Sections 1 and 2."),
 ("The ablation is only on Mine-Objects and lacks evidence of generalizability.", "We partially agree and added a compact five-row component ablation on ExDark validation. The RF-DETR baseline obtains 0.518 AP@0.5:0.95 and the full model obtains 0.563; the held-out ExDark comparison is reported separately. We did not repeat the complete ablation on ScienceDB because its converted validation partition is used as an additional mine-domain benchmark rather than as a second full development suite. This scope decision is now stated explicitly, and no universal cross-dataset claim is made (Section 4.2 and Table 4)."),
 ("It is unclear whether YOLO baselines were tuned.", "We clarified the baseline protocol in Section 4.1. YOLO baselines use official model definitions, the common input resolution and 75-epoch budget, and recipe-level checks on the validation split for learning rate, batch size, and augmentation strength. We do not claim an exhaustive architecture-specific hyperparameter search; instead, we report the reproduced protocol and state this limitation explicitly. The controlled RF-DETR ablations use identical settings across rows."),
 ("FFT/IFFT reduces FPS by approximately 18% without sufficient deployment discussion.", "We expanded the Efficiency Analysis and Limitations. Under the stated batch-1, 560x560, RTX 3090 protocol, throughput decreases from 21.3 to 17.5 FPS and latency increases from approximately 46.9 to 57.1 ms per image. We now state that this does not meet a 25--30 FPS target on the reported hardware but is compatible with moderate-rate monitoring. DQCD and SCU are training-only; selected-tap deployment and spatial-domain alternatives are identified as future efficiency studies rather than unsupported claims."),
 ("Figure 5 lacks an intuitive baseline comparison.", "Figure 5 already contained the RF-DETR baseline in the middle row for the same input shown in the top row. We agree that the original caption described the arrangement ambiguously. We revised the caption to state explicitly that each column compares the input, RF-DETR prediction, and Spectral-DETR prediction on exactly the same image. The examples remain qualitative illustrations and are not used as quantitative evidence."),
 ("Mine-Objects is small and has a higher overfitting risk.", "We agree that the dataset is compact and now state this as a limitation. The split is performed at acquisition-sequence level before frame selection, so adjacent frames from one source sequence cannot cross partitions; the 8:1:1 split contains 2,464 training, 308 validation, and 309 test images. Checkpoints and ablation decisions use validation only, while the held-out test set is reserved for the final comparison. The ExDark and ScienceDB evaluations provide complementary external-domain evidence, but we do not claim universal mine-domain generalization."),
 ("The Pearson correlation of 0.60 is only moderate and should not be treated as key evidence.", "We agree and have narrowed the interpretation. The Pearson value of 0.60, together with Spearman 0.58, is now reported only as a supporting diagnostic that predicted log-variance ranks part of localization difficulty. It is not presented as conclusive uncertainty calibration or as a physical degradation measurement. The main SCU evidence is instead the controlled ablation, AP_S improvement, and coefficient-sensitivity analysis (Section 5.3)."),
],
2: [
 ("Please do not use abbreviations such as P@0.5 in the abstract.", "We removed the ambiguous P@0.5 notation from the abstract and now spell out average precision at the IoU threshold of 0.5 and the averaged 0.5--0.95 IoU range on first presentation. The abbreviation is retained only in the technical sections and table headings where the metric is defined."),
 ("Please add numbers identifying equations.", "All displayed mathematical expressions use numbered equation environments in the revised manuscript. The temperature and SCU equations additionally carry labels for direct cross-reference, including Equation (\ref{eq:tau}) and Equation (\ref{eq:scu_target})."),
 ("What is FiLM modulation in Figure 2?", "We added a definition in Section 3.1. FiLM denotes Feature-wise Linear Modulation: a scene encoder predicts a channel-wise scale gamma and shift beta from the log-amplitude spectrum, and these parameters modulate the raw gate logits before the sigmoid gate. Figure 2 and its caption now identify this operation explicitly."),
 ("What is a 2D Real FFT?", "We now state that the 2D real FFT applies a two-dimensional FFT to a real-valued feature map. Conjugate symmetry allows the one-sided representation with W_fft = floor(W_f/2)+1 coefficients on the last frequency axis, followed by the corresponding inverse real FFT."),
 ("What are B, H, and W?", "The notation paragraph now defines B as batch size, H as image height, and W as image width. Feature-map dimensions are separately denoted by C, H_f, and W_f."),
 ("What are u and v?", "We define u and v as the discrete vertical and horizontal frequency indices. Their normalized signed frequencies are xi_u and eta_v, generated by fftfreq and rfftfreq, respectively; the radial frequency is r(u,v)=sqrt(xi_u^2+eta_v^2)."),
 ("Please explain the dot-inside-circle operator.", "The revised text states that the circled-dot symbol \odot denotes element-wise multiplication with broadcasting over batch and channel dimensions. It is distinct from the ordinary centered dot used for scalar multiplication."),
 ("What is E[G_k] and what do the square brackets mean?", "We now explain that E[G_k] denotes the arithmetic mean of the gate values over the batch, channel, and frequency dimensions. The brackets are expectation notation, not an additional operator or indexing convention."),
 ("What does the dot operator mean?", "We clarified the notation: a centered dot denotes ordinary scalar multiplication, whereas h_i^T h_n denotes the vector inner product. Element-wise tensor multiplication is written with \odot and is defined separately."),
 ("Please specify exactly which data were used.", "The dataset section now identifies all three sources precisely. Mine-Objects contains 3,081 images and 14 categories with a sequence-level 8:1:1 split. ScienceDB is the V1 Coal Mine Underground Drilling Site Object Detection Dataset, DOI 10.57760/sciencedb.j00001.01020, containing 70,948 images and five categories; we report its deterministic converted validation partition. ExDark contains 7,363 images, 12 categories, and the dataset-specific train/validation/test evaluation described in Section 4.1."),
 ("The source code could not be viewed at the ScienceDB URL and returned 404.", "We clarified that ScienceDB is the data record, not the software repository. The Data Availability statement now provides the exact ScienceDB DOI for the dataset and a separate public GitHub URL for the Spectral-DETR source code. The code repository is intended to contain the model implementation, conversion scripts, evaluation tools, and the final configuration used for the reported protocol."),
],
3: [
 ("Why are three fixed Gaussian bands sufficient, and is dynamic allocation explored?", "We added a DAFD-only band-count analysis for one through five bands. AP@0.5:0.95 rises from 0.472 for one band to 0.475 for three bands; four bands ties on the strict metric but does not improve AP@0.5 or AP_S. We therefore describe three bands as the best overall trade-off rather than a unique optimum. Per-image dynamic band allocation is retained as future work because the present experiments do not validate an additional controller."),
 ("The gate statistic is not shown to correlate with ground-truth degradation severity.", "We agree that the gate mean is not a calibrated physical severity score. The revised manuscript explicitly states this limitation. Figure 3 and the degradation-stratified results are descriptive, while the fixed, shuffled, and random-gate controls test whether the signal is useful for the DQCD interface beyond adding modules independently. We consequently use cautious language: the controls support a modest coupling benefit, not a claim that the gate statistic recovers ground-truth severity across degradation types."),
 ("SCU coefficients may be dataset-dependent, but sensitivity is not analyzed.", "We added a coefficient-sensitivity table covering the default slope and center as well as perturbed values. AP@0.5:0.95 varies by no more than 0.001 across these local settings. We now describe the coefficients as a soft geometric prior and explicitly acknowledge that optimal values may depend on object-size distributions outside the tested range."),
 ("The cascaded modules are not tested for synergy or removal.", "The primary Mine-Objects validation ablation now includes all single-module, pairwise, and full configurations. The full gain is 0.014 versus an individual-gain sum of 0.011, which we describe as modest non-additivity rather than strong synergy. The DAFD-to-DQCD fixed, shuffled, and random controls further test the explicit coupling interface. These experiments also show the performance cost of removing any stage without requiring a second full architecture search."),
 ("Recent degraded DETR, learnable-filter, and multimodal detection work is missing.", "We expanded Section 2 with recent low-light and degraded-scene DETR variants, including MPE-DETR, MDFD2-DETR, and frequency-enhanced Transformers, and clarified that frequency processing itself is not claimed as novel. The suggested multimodal road-defect and saliency methods address different sensing modalities or task objectives; we did not present them as directly comparable detector baselines, but we acknowledge this boundary and position Spectral-DETR by its detector-internal cross-stage interface."),
 ("The category-level Sm-AP definition and its limitations are unclear.", "We introduced the Sm-AP definition before the benchmark tables and explicitly distinguish it from instance-level COCO AP_S. Sm-AP selects categories whose mean instance area is below 32^2 pixels to provide a stable category-level diagnostic across the small Mine-Objects and ExDark benchmarks. We also added the limitation that category means can obscure substantial within-class scale variation; Sm-AP is therefore complementary and is not used as a replacement for COCO AP_S."),
 ("The ScienceDB AP@0.5 and AP@0.5:0.95 gap is very large.", "We added a dedicated interpretation. The gap is 0.478, indicating strong coarse-threshold detection but weaker localization at stricter IoUs. The same qualitative pattern appears for all listed detectors. We discuss small or elongated targets, ambiguous boundaries, and annotation variability as plausible contributors, while noting that AP@0.75 was not measured and therefore the gap cannot be decomposed into specific IoU regimes."),
 ("The FFT overhead is not compared across selected feature levels or against spatial alternatives.", "We expanded the efficiency limitations to state that DAFD can be applied to selected encoder taps and that tap-specific accuracy/latency measurements and a matched spatial-domain comparison are future work. We did not add an unvalidated alternative experiment to the current revision; the reported 21.3-to-17.5 FPS measurement is retained as the hardware-specific efficiency result."),
],
4: [
 ("The literature review should use coherent paragraphs rather than list format.", "Section 2 was rewritten as connected narrative paragraphs. The detector, degraded-scene, contrastive-learning, and uncertainty-regression literature are now synthesized in relation to the three stages rather than presented as a list of disconnected methods."),
 ("The methodological transparency and fixed 8:1:1 split are reasonable.", "We retained the fixed sequence-level 8:1:1 split and expanded the dataset and evaluation-protocol descriptions to state the exact training, validation, and test roles. The limitation of a compact dataset and rare classes is now discussed explicitly."),
 ("Raw Gaussian-mask center and width parameters should be italicized.", "The center and width parameters are now written as mathematical variables, with bold symbols used only for parameter vectors and italic symbols for individual scalar components in the mask equations."),
 ("Check the piecewise formatting for the temperature equation.", "The temperature equation was checked in the revised LaTeX source and retains a numbered, aligned cases expression. The surrounding text defines the active-DAFD and inactive-DAFD conditions and gives the base temperature."),
 ("Introduce Sm-AP earlier in Section 4.2.", "The Sm-AP definition now appears in the Evaluation Protocol before the Mine-Objects and ExDark benchmark tables. The table captions repeat the category subsets and distinguish Sm-AP from COCO AP_S."),
 ("Move limitations from bullet/run-in format into standard paragraphs.", "The Limitations and Future Work section is now written as three continuous paragraphs covering computational cost, design sensitivity/generalization, and statistical/data limitations. No bullet-style run-in headings are used."),
 ("Check citation numbering and journal abbreviations.", "We performed a local citation audit: all in-text citation keys resolve to the bibliography, there are no duplicate keys or placeholder citations, and the recent DOI records are populated. The final compiled PDF will be used for the last MDPI bracket-numbering and journal-name style check."),
],
}


def clear_paragraph(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_cell_text(cell, text, color=None, bold=False, size=9.5):
    # Keep the template cell geometry and borders, but replace its content.
    first = cell.paragraphs[0]
    for p in list(cell.paragraphs[1:]):
        p._element.getparent().remove(p._element)
    clear_paragraph(first)
    first.alignment = WD_ALIGN_PARAGRAPH.LEFT
    first.paragraph_format.space_after = Pt(4)
    run = first.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_paragraphs(cell, paragraphs, color=(128, 0, 0)):
    first = cell.paragraphs[0]
    for p in list(cell.paragraphs[1:]):
        p._element.getparent().remove(p._element)
    clear_paragraph(first)
    for idx, text in enumerate(paragraphs):
        p = first if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(*color)


def cell_text(cell):
    return " ".join(" ".join(p.text.split()) for p in cell.paragraphs).strip()


def row_label(table, row_idx):
    return cell_text(table.rows[row_idx].cells[0])


def prepare_document(reviewer):
    doc = Document(SOURCE[reviewer])
    table = doc.tables[0]

    # The template uses merged cells; editing the first physical cell updates the merged content.
    set_cell_text(table.cell(0, 0), f"Response to Reviewer {reviewer} Comments", bold=True, size=16)
    set_cell_text(table.cell(2, 0), SUMMARY, color=(0, 0, 0), size=10)

    general_start = next(
        idx for idx in range(len(table.rows))
        if row_label(table, idx).startswith("Does the introduction provide")
    )
    for row_idx, (_, evaluation, response) in enumerate(GENERAL[reviewer], start=general_start):
        set_cell_text(table.cell(row_idx, 1), evaluation, color=(0, 0, 0), size=9.5)
        set_cell_text(table.cell(row_idx, 2), response, color=(0, 76, 153), size=9.5)

    comment_rows = [
        idx for idx in range(len(table.rows))
        if row_label(table, idx).lower().startswith("comments ")
    ]
    if len(comment_rows) != len(COMMENTS[reviewer]):
        raise ValueError(
            f"Reviewer {reviewer}: template has {len(comment_rows)} comment rows, "
            f"but {len(COMMENTS[reviewer])} responses are configured"
        )
    for row_idx, (_, response) in zip(comment_rows, COMMENTS[reviewer]):
        response_row = row_idx + 1
        set_cell_paragraphs(table.cell(response_row, 2), ["Response and revision: " + response])

    # Fill the language section and remove its template-only response placeholder.
    language_header = next(
        idx for idx in range(len(table.rows))
        if row_label(table, idx).startswith("4. Response to Comments on the Quality")
    )
    language_point = language_header + 1
    language_response = language_header + 2
    set_cell_text(
        table.cell(language_response, 2),
        "We performed a focused language revision for clarity, grammar, terminology consistency, and paragraph flow while preserving the technical content.",
        color=(0, 76, 153),
        size=9.5,
    )
    clarification_header = next(
        idx for idx in range(len(table.rows))
        if row_label(table, idx).startswith("5. Additional clarifications")
    )
    clarification_row = clarification_header + 1
    set_cell_text(
        table.cell(clarification_row, 0),
        "No additional clarifications beyond the point-by-point responses above.",
        color=(0, 0, 0),
        size=9.5,
    )

    # Replace the red instructional text in the summary row if it survived a merged-cell alias.
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if "[This is only a recommended summary" in run.text or "[Here, mention" in run.text:
                        run.text = ""

    # Keep the source's page geometry and table style; update core document properties.
    doc.core_properties.title = f"Response to Reviewer {reviewer} Comments"
    doc.core_properties.subject = "Point-by-point response for revised Spectral-DETR manuscript"
    return doc


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for reviewer in range(1, 5):
        output = OUT_DIR / f"Response_to_Reviewer_{reviewer}.docx"
        prepare_document(reviewer).save(output)
        print(output)


if __name__ == "__main__":
    main()
